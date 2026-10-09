"""Bounded CSV/XLSX records, reviewed column mappings and atomic append. Never execute uploaded SQL."""
import csv
import calendar
import hashlib
import io
import json
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from typing import Literal
from zipfile import BadZipFile, ZipFile

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field

from .analytics import identifier, now
from .database import connect, metadata
from .forms import CustomerFields
from .operations import owner
from .security import actor, require

router = APIRouter(prefix='/data/imports')
Entity = Literal['customers', 'accounts', 'transactions', 'loans']
MAX_BYTES = 5 * 1024 * 1024


def field(key, label, required=True, hint=''):
    return dict(key=key, label=label, required=required, hint=hint)


SCHEMAS = {
    'customers': [field('id', 'Customer ID', hint='Unique CUS- ID. Keep the same ID in every linked sheet.'),
                  field('fullName', 'Full name'), field('dateOfBirth', 'Date of birth', hint='YYYY-MM-DD'),
                  *[field(key, label, False) for key, label in [('email','Email'),('phone','Phone'),('address','Address'),('city','City'),('state','State'),('postalCode','Postal code')]]],
    'accounts': [field('id','Account ID'), field('customer_id','Customer ID', hint='Import the customer first.'), field('opened_on','Opening date', hint='YYYY-MM-DD')],
    'transactions': [field('id','Transaction ID'), field('account_id','Account ID'), field('occurred_at','Transaction timestamp',hint='ISO timestamp. Unzoned Excel dates use the timezone selected below.'),
                     field('amount','Amount',hint='Decimal currency units, not paise/cents.'), field('currency','Currency'), field('region','Region'), field('merchant','Merchant'), field('channel','Channel'),
                     field('device_id','Device ID',False), field('new_device','New device',hint='0 or 1'), field('flagged','Source fraud flag',hint='0 or 1; an unverified source observation.'), field('score','Source risk score',False,'0–100 if supplied. Blank remains unknown, not zero.')],
    'loans': [field('id','Loan ID'), field('customer_id','Customer ID'), field('branch_id','Branch ID'), field('loan_type','Loan type'), field('principal','Original principal',hint='Decimal currency units.'),
              field('default_months','Reported default months',False,'Semicolon-separated YYYY-MM values. Blank means no supplied default observations.')],
}
TABLES = {'customers':'customers', 'accounts':'accounts', 'transactions':'transactions', 'loans':'loans'}


class ReviewInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    uploadId: str = Field(pattern=r'^UPLOAD-[a-zA-Z0-9]+$')
    entity: Entity
    sheet: str = Field(min_length=1,max_length=100)
    mapping: dict[str,str] = Field(max_length=40)
    timezone: Literal['+05:30','+00:00'] = '+05:30'


class ApplyInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    previewId: str = Field(pattern=r'^REVIEW-[a-zA-Z0-9]+$')
    sourceHash: str = Field(pattern=r'^[a-f0-9]{64}$')
    confirmed: Literal[True]


def digest(value):
    return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False).encode()).hexdigest()


def cell_value(value):
    if value is None: return ''
    if isinstance(value,(datetime,date)): return value.isoformat()
    if isinstance(value,bool): return '1' if value else '0'
    result = str(value).strip()
    if len(result)>2000: raise ValueError('A cell exceeds 2,000 characters. Import only record fields.')
    if result.startswith('='): raise ValueError('Formula cells are not supported. Export or paste values only.')
    return result


def collect_sheet(name, source):
    values = []
    for row_number,row in enumerate(source,1):
        # Bounded at the worksheet parser too, including formatted empty trailing cells.
        if len(row)>40: raise ValueError('Use at most 40 columns per sheet.')
        row = [cell_value(cell) for cell in row]
        if not any(row): continue
        if row_number>2001: raise ValueError('Remove unused rows. Data must be within the first 2,000 physical rows.')
        if not values:
            while row and not row[-1]: row.pop()
            if not row or any(not h for h in row) or len(set(row))!=len(row):
                raise ValueError(f'{name}: use a header row with unique, non-empty column names.')
            values.append((row_number,row))
            continue
        width = len(values[0][1])
        if len(row)>width and any(row[width:]): raise ValueError(f'{name}: row {row_number} has values without column headers.')
        values.append((row_number,(row+['']*width)[:width]))
        if len(values)>501: raise ValueError('Import at most 500 records per sheet. Split larger exports into batches.')
    if not values: return None
    headers = values[0][1]
    return dict(name=name,headers=headers,rows=[dict(row=n,values=dict(zip(headers,row))) for n,row in values[1:]])


def parse_file(content, filename):
    suffix = PurePath(filename).suffix.lower()
    if suffix not in ('.csv','.xlsx'):
        raise ValueError('Use .xlsx or UTF-8 .csv. For SQL databases, export table records to CSV/Excel; .sql commands are not executed.')
    if not content or len(content)>MAX_BYTES: raise ValueError('Use a non-empty file up to 5 MB.')
    if suffix=='.csv':
        try:
            text=content.decode('utf-8-sig')
            sample=text[:8192]
            try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
            except csv.Error: dialect=csv.excel
            sheet=collect_sheet('CSV records',csv.reader(io.StringIO(text),dialect,strict=True))
        except (UnicodeError,csv.Error) as error:
            raise ValueError('Use a readable UTF-8 CSV with a header row.') from error
        sheets=[sheet] if sheet else []
    else:
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                members=archive.infolist()
                if len(members)>1000 or sum(m.file_size for m in members)>20*1024*1024:
                    raise ValueError('The workbook expands beyond the 20 MB processing limit. Use a smaller export.')
                if any('vba' in m.filename.lower() or '/externallinks/' in m.filename.lower() for m in members):
                    raise ValueError('Use a workbook without macros or external workbook links.')
                from defusedxml.ElementTree import iterparse
                for member in members:
                    if member.filename.startswith('xl/worksheets/') and member.filename.endswith('.xml'):
                        with archive.open(member) as stream:
                            for _,element in iterparse(stream,events=('end',)):
                                if element.tag.endswith('}c'):
                                    match=re.fullmatch(r'([A-Z]+)(\d+)',element.get('r',''))
                                    if not match: raise ValueError('Workbook has an invalid cell address.')
                                    column=0
                                    for letter in match[1]: column=column*26+ord(letter)-64
                                    if column>40 or int(match[2])>2001: raise ValueError('Keep worksheet content within 40 columns and 2,000 physical rows. Remove unused formatting before exporting.')
                                element.clear()
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
            try:
                if len(book.worksheets)>8: raise ValueError('Use at most 8 worksheets per workbook.')
                sheets=[]
                for sheet in book.worksheets:
                    # Do not trust workbook dimensions; read actual cells with explicit bounds.
                    sheet.reset_dimensions()
                    def rows():
                        for cells in sheet.iter_rows(max_row=2002):
                            if any(c.data_type in ('f','e') for c in cells):
                                raise ValueError(f'{sheet.title}: formulas or Excel errors found. Paste values before importing.')
                            values=[c.value for c in cells]
                            while values and values[-1] is None: values.pop()
                            yield values
                    result=collect_sheet(sheet.title,rows())
                    if result: sheets.append(result)
            finally: book.close()
        except ValueError:
            raise
        except Exception as error:
            raise ValueError('This is not a readable .xlsx workbook. Re-save it as Excel Workbook (.xlsx).') from error
    if not sheets or not any(s['rows'] for s in sheets): raise ValueError('Add at least one data row below the headers.')
    if sum(len(s['rows']) for s in sheets)>2000 or len(json.dumps(sheets).encode())>2*1024*1024:
        raise ValueError('Use at most 2,000 total records and 2 MB of cell content per upload.')
    return suffix[1:],sheets


@router.get('/schema')
def schemas(datasetId: Literal['meridian','bankshield']='meridian'):
    meta=metadata(datasetId)
    with connect() as db:
        branches=[dict(id=r['branch_id'],name=r['branch_name']) for r in db.execute('SELECT DISTINCT branch_id,branch_name FROM customer_master WHERE dataset=? ORDER BY branch_id',(datasetId,))]
    return dict(entities=SCHEMAS if datasetId=='meridian' else {'transactions':SCHEMAS['transactions']},periods=meta['periods'],regions=meta.get('regions',[]),branches=branches,currency=meta['currency'])


@router.post('/file')
def upload(file: UploadFile=File(...),datasetId: Literal['meridian','bankshield']=Form('meridian')):
    require('analyze'); metadata(datasetId)
    filename=(file.filename or 'upload').replace('\\','/').split('/')[-1][:180]
    try: content=file.file.read(MAX_BYTES+1)
    finally: file.file.close()
    kind,sheets=parse_file(content,filename)
    upload_id=identifier('UPLOAD')
    with connect() as db:
        db.execute('INSERT INTO file_import_uploads VALUES(?,?,?,?,?,?,?,?)',(upload_id,owner(),datasetId,filename,kind,hashlib.sha256(content).hexdigest(),json.dumps(sheets),now()))
    return dict(id=upload_id,filename=filename,format=kind,datasetId=datasetId,sheets=[dict(name=s['name'],headers=s['headers'],rows=len(s['rows']),sample=[r['values'] for r in s['rows'][:3]]) for s in sheets])


def owned_upload(db,upload_id,identity):
    row=db.execute('SELECT * FROM file_import_uploads WHERE id=?',(upload_id,)).fetchone()
    if not row or row['owner']!=identity: raise PermissionError('This upload does not belong to this workspace/staff actor.')
    return row


def money(value):
    try:
        number=Decimal(value)
        if not number.is_finite() or not 0<number<=1000000000 or number*100!=(number*100).to_integral_value(): raise ValueError()
        return int(number*100)
    except (InvalidOperation,ValueError): raise ValueError('Use a positive amount with at most two decimals.')


def calendar_date(value):
    if 'T' in value or ' ' in value:
        return datetime.fromisoformat(value).date().isoformat()
    return date.fromisoformat(value).isoformat()


def validate_rows(db,dataset,entity,rows):
    if dataset!='meridian' and entity!='transactions': raise ValueError('BankShield accepts only transaction appends; its customer relationships are incomplete.')
    meta=json.loads(db.execute('SELECT metadata FROM datasets WHERE id=?',(dataset,)).fetchone()[0])
    errors=[];valid=[];seen=set();seen_people=set()
    for item in rows:
        row=item['values']; reasons=[]; record_id=row.get('id','')
        if not re.fullmatch(r'[a-zA-Z0-9-]{1,80}',record_id): reasons.append('Use a unique record ID containing letters, numbers or hyphens.')
        if record_id in seen: reasons.append('Duplicate ID within the selected sheet.')
        seen.add(record_id)
        if db.execute(f'SELECT 1 FROM {TABLES[entity]} WHERE dataset=? AND id=?',(dataset,record_id)).fetchone(): reasons.append('Record ID already exists. Imports append new records; they do not overwrite.')
        if db.execute('SELECT 1 FROM record_import_sources WHERE dataset=? AND entity=? AND record_id=?',(dataset,entity,record_id)).fetchone(): reasons.append('This ID has an existing import receipt. Use a new ID; retained import history cannot be overwritten.')
        try:
            if entity=='customers':
                if not re.fullmatch(r'CUS-[a-zA-Z0-9-]{1,76}',record_id): raise ValueError('Customer IDs must start with CUS- to support directory navigation.')
                if db.execute('SELECT 1 FROM customer_profiles WHERE id=?',(record_id,)).fetchone(): raise ValueError('Customer ID is already used by another profile.')
                candidate={f['key']:row.get(f['key'],'') for f in SCHEMAS[entity] if f['key']!='id'}
                candidate['dateOfBirth']=calendar_date(candidate['dateOfBirth'])
                values=CustomerFields.model_validate(candidate).model_dump()
                if len(values['fullName'].strip())<2 or not values['dateOfBirth']: raise ValueError('Supply the full name and date of birth.')
                # Conservative identity review; never auto-merge people.
                from .form_quality import duplicates,normalized
                if duplicates(db,values): raise ValueError('Possible existing customer. Review their profile manually instead of importing a duplicate.')
                identity_keys={(values['dateOfBirth'],'name',normalized(values['fullName']))}
                if values['email']: identity_keys.add((values['dateOfBirth'],'email',normalized(values['email'])))
                if values['phone']: identity_keys.add((values['dateOfBirth'],'phone',re.sub(r'\D','',values['phone'])))
                if seen_people & identity_keys: raise ValueError('Possible duplicate person within this sheet. Review the rows manually before importing.')
                seen_people.update(identity_keys)
                result=(dataset,record_id,values)
            elif entity=='accounts':
                if not db.execute('SELECT 1 FROM customers WHERE dataset=? AND id=?',(dataset,row['customer_id'])).fetchone(): raise ValueError('Customer parent is missing. Import customers first.')
                opened=date.fromisoformat(calendar_date(row['opened_on']))
                if opened>date.today() or opened.year<1900: raise ValueError('Use a valid account opening date, not a future date.')
                result=(dataset,record_id,row['customer_id'],opened.isoformat())
            elif entity=='loans':
                if not db.execute('SELECT 1 FROM customers WHERE dataset=? AND id=?',(dataset,row['customer_id'])).fetchone(): raise ValueError('Customer parent is missing. Import customers first.')
                if not db.execute('SELECT 1 FROM customer_master WHERE dataset=? AND branch_id=?',(dataset,row['branch_id'])).fetchone(): raise ValueError('Choose an existing branch ID from the dataset.')
                parent=db.execute('SELECT branch_id FROM customer_master WHERE dataset=? AND id=?',(dataset,row['customer_id'])).fetchone()
                if parent and parent['branch_id']!=row['branch_id']: raise ValueError('Loan branch must match this historical customer branch.')
                if not row['loan_type'] or len(row['loan_type'])>80: raise ValueError('Supply a loan type up to 80 characters.')
                months=[v.strip() for v in row.get('default_months','').split(';') if v.strip()]
                if any(month not in meta['periods'] for month in months) or len(set(months))!=len(months): raise ValueError('Default months must be unique supported YYYY-MM periods separated by semicolons.')
                result=(dataset,record_id,row['customer_id'],row['branch_id'],row['loan_type'],money(row['principal']),json.dumps([calendar.month_abbr[int(month[-2:])] for month in months]))
            else:
                from .data_quality import COLUMNS,validate
                tx={key:row.get(key,'') for key in COLUMNS}
                score_missing=not tx['score']
                if score_missing: tx['score']='0' # Validate the other fields; persist NULL below.
                tx_valid,tx_errors=validate(db,dataset,[tx])
                if tx_errors: reasons.extend(tx_errors[0]['reasons']); result=None
                else: result=(*tx_valid[0][:-1],None) if score_missing else tx_valid[0]
            if not reasons: valid.append(result)
        except (ValueError,KeyError) as error:
            if isinstance(error,KeyError): reasons.append(f'Missing required field: {error.args[0]}.')
            else: reasons.append(str(error))
        if reasons: errors.append(dict(row=item['row'],reasons=reasons))
    return valid,errors


@router.post('/review')
def review(body: ReviewInput):
    require('analyze'); identity=owner()
    with connect() as db:
        upload=owned_upload(db,body.uploadId,identity)
        sheet=next((s for s in json.loads(upload['sheets']) if s['name']==body.sheet),None)
        if not sheet or not sheet['rows']: raise ValueError('Choose a sheet with records.')
        fields=SCHEMAS[body.entity]; allowed={f['key'] for f in fields}
        mapping={key:value for key,value in body.mapping.items() if value}
        if set(mapping)-allowed or any(h not in sheet['headers'] for h in mapping.values()): raise ValueError('Map only the displayed record fields and uploaded columns.')
        if len(set(mapping.values()))!=len(mapping): raise ValueError('Each uploaded column can map to only one record field.')
        missing=[f['label'] for f in fields if f['required'] and f['key'] not in mapping]
        if missing: raise ValueError('Map required fields: '+', '.join(missing)+'.')
        rows=[]
        for source in sheet['rows']:
            values={f['key']:source['values'].get(mapping.get(f['key'],''),'') for f in fields}
            if body.entity=='transactions':
                try:
                    stamp=datetime.fromisoformat(values['occurred_at'].replace('Z','+00:00'))
                    if stamp.tzinfo is None:
                        stamp=stamp.replace(tzinfo=timezone(timedelta(hours=5,minutes=30)) if body.timezone=='+05:30' else timezone.utc)
                    values['occurred_at']=stamp.isoformat()
                except ValueError: pass # Validator reports the original row instead of aborting preview.
            rows.append(dict(row=source['row'],values=values))
        valid,errors=validate_rows(db,upload['dataset'],body.entity,rows)
        meta=json.loads(db.execute('SELECT metadata FROM datasets WHERE id=?',(upload['dataset'],)).fetchone()[0])
        preview_id=identifier('REVIEW');source_hash=digest(dict(uploadHash=upload['source_hash'],entity=body.entity,sheet=body.sheet,mapping=mapping,timezone=body.timezone,rows=rows))
        db.execute('INSERT INTO file_import_reviews VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(preview_id,upload['id'],identity,body.entity,body.sheet,json.dumps(mapping),json.dumps(rows),json.dumps(errors),meta['version'],source_hash,now(),'preview',None))
    return dict(id=preview_id,sourceHash=source_hash,datasetId=upload['dataset'],entity=body.entity,currency=meta['currency'],rows=len(rows),validRows=len(valid),canImport=not errors,errors=errors,sample=[r['values'] for r in rows[:5]])


@router.post('/apply')
def apply(body: ApplyInput):
    require('customer_write'); identity=owner()
    with connect() as db:
        db.execute('BEGIN IMMEDIATE')
        review=db.execute('SELECT * FROM file_import_reviews WHERE id=?',(body.previewId,)).fetchone()
        if not review or review['owner']!=identity: raise PermissionError('This review does not belong to this workspace/staff actor.')
        upload=owned_upload(db,review['upload_id'],identity)
        if review['source_hash']!=body.sourceHash: raise ValueError('The confirmed preview does not match the reviewed content.')
        if review['status']=='imported': return {**json.loads(review['receipt']),'replayed':True}
        dataset=upload['dataset'];entity=review['entity'];rows=json.loads(review['rows'])
        meta=json.loads(db.execute('SELECT metadata FROM datasets WHERE id=?',(dataset,)).fetchone()[0])
        if meta['version']!=review['base_version']: raise HTTPException(409,'Dataset changed since preview. Validate again before confirming.')
        valid,errors=validate_rows(db,dataset,entity,rows)
        if errors: raise ValueError('Import blocked: records or parent links changed. Validate again; no rows were saved.')
        if entity=='customers':
            for _,record_id,values in valid:
                db.execute('INSERT INTO customers VALUES(?,?)',(dataset,record_id))
                db.execute('INSERT INTO customer_profiles VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(record_id,dataset,*[values[f['key']] for f in SCHEMAS['customers'] if f['key']!='id'],'pending_verification',None,now()))
        else:
            placeholders=','.join('?' for _ in valid[0])
            db.executemany(f'INSERT INTO {TABLES[entity]} VALUES({placeholders})',valid)
            if entity=='transactions': db.executemany('INSERT INTO device_events VALUES(?,?,?,?)',[(r[0],r[1],r[9],r[3]) for r in valid if r[9]])
        db.executemany('INSERT INTO record_import_sources VALUES(?,?,?,?)',[(dataset,entity,r['values']['id'],review['id']) for r in rows])
        meta.setdefault('baseFixtureVersion',meta['version'])
        meta['version']=digest([meta['version'],review['source_hash']])
        meta['coverage'][entity]=db.execute(f'SELECT COUNT(*) FROM {TABLES[entity]} WHERE dataset=?',(dataset,)).fetchone()[0]
        meta['coverage']['completeCustomerChains']=db.execute('SELECT COUNT(*) FROM transactions t JOIN accounts a ON a.dataset=t.dataset AND a.id=t.account_id JOIN customers c ON c.dataset=a.dataset AND c.id=a.customer_id WHERE t.dataset=?',(dataset,)).fetchone()[0]
        note='Reviewed CSV/Excel appends present. Source flags, scores and default markers are unverified observations; imported loans do not invent repayment schedules.'
        if note not in meta['limitations']: meta['limitations'].append(note)
        db.execute('UPDATE datasets SET metadata=? WHERE id=?',(json.dumps(meta),dataset))
        receipt=dict(previewId=review['id'],entity=entity,importedRows=len(valid),datasetVersion=meta['version'],replayed=False)
        db.execute("UPDATE file_import_reviews SET status='imported',receipt=? WHERE id=?",(json.dumps(receipt),review['id']))
        db.execute('INSERT INTO audit_runs VALUES(?,?,?,?,?,?)',(identifier('RUN'),'reviewed_file_import',json.dumps(dict(datasetId=dataset,uploadId=upload['id'],previewId=review['id'],entity=entity,actor=actor().actor_id)),'Explicit confirmation of reviewed record append',now(),review['source_hash']))
    return receipt


@router.get('/history')
def history(datasetId: Literal['meridian','bankshield']='meridian'):
    with connect() as db:
        rows=db.execute('SELECT r.id,r.entity,r.sheet,r.status,r.created_at,r.receipt,r.errors,r.rows,u.filename,u.format FROM file_import_reviews r JOIN file_import_uploads u ON u.id=r.upload_id WHERE r.owner=? AND u.dataset=? ORDER BY r.created_at DESC,r.id DESC LIMIT 20',(owner(),datasetId)).fetchall()
    return [dict(id=r['id'],filename=r['filename'],format=r['format'],entity=r['entity'],sheet=r['sheet'],status=r['status'],createdAt=r['created_at'],rows=len(json.loads(r['rows'])),errorRows=len(json.loads(r['errors'])),importedRows=json.loads(r['receipt'])['importedRows'] if r['receipt'] else 0) for r in rows]


@router.get('/template')
def template(entity: Entity='customers',datasetId: Literal['meridian','bankshield']='meridian'):
    meta=metadata(datasetId)
    if datasetId=='bankshield' and entity!='transactions': raise ValueError('BankShield supports transaction templates only.')
    headers=[f['key'] for f in SCHEMAS[entity]]
    samples={
        'customers':['CUS-IMPORT-001','Asha Demo','1992-06-15','asha@example.test','9876543210','12 Demo Road','Ahmedabad','Gujarat','380001'],
        'accounts':['ACC-IMPORT-001','CUS-IMPORT-001','2026-01-01'],
        'transactions':['TX-IMPORT-001','ACC-IMPORT-001','2026-09-15T10:30:00+05:30','1500.00','INR','Ahmedabad','Demo Retail','Online','','0','0','10'],
        'loans':['LN-IMPORT-001','CUS-IMPORT-001','BR-AHM','Personal','250000.00',''],
    }
    if datasetId=='bankshield':
        with connect() as db:
            source=db.execute('SELECT t.* FROM transactions t JOIN accounts a ON a.dataset=t.dataset AND a.id=t.account_id JOIN customers c ON c.dataset=a.dataset AND c.id=a.customer_id WHERE t.dataset=? AND t.occurred_at IS NOT NULL ORDER BY t.occurred_at DESC LIMIT 1',(datasetId,)).fetchone()
        if not source: raise ValueError('No linked source transaction is available for a template.')
        samples['transactions']=['TX-IMPORT-001',source['account_id'],source['occurred_at'],str(Decimal(source['amount_minor'])/100),meta['currency'],source['region'] or 'Unknown',source['merchant'] or 'Demo merchant',source['channel'] or 'Unknown','','0',str(source['flagged']),str(source['score']) if source['score'] is not None else '']
    output=io.StringIO();writer=csv.writer(output);writer.writerow(headers);writer.writerow(samples[entity])
    return Response(output.getvalue(),media_type='text/csv',headers={'Content-Disposition':f'attachment; filename="bankweave-{entity}-template.csv"'})
